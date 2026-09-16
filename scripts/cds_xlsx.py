#!/usr/bin/env python3
"""
cds_xlsx.py — Flatten a Common Data Set Excel workbook into plain text.

CDS xlsx files are already tabular (one sheet per section, each row a
label/value pair), so there's no PDF-style layout ambiguity to fight —
reading cell values directly and handing Claude a flat text dump is both
cheaper (no document/vision tokens) and more reliable than converting to
PDF and reusing the native-document path.
"""

from pathlib import Path

# Sheets that carry no CDS data — the definitions appendix (already excluded
# from the PDF path via SYSTEM_PROMPT) and administrative cover sheets some
# schools include in their CDS workbook.
SKIP_SHEETS = {"cds definitions", "welcome", "answer sheet"}


def xlsx_to_text(xlsx_path: Path) -> str:
    """Flatten every data sheet into '=== SHEET: name ===' + 'label | value' lines."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    chunks = []
    for name in wb.sheetnames:
        if name.strip().lower() in SKIP_SHEETS:
            continue
        ws = wb[name]
        rows = []
        for row in ws.iter_rows():
            vals = [str(c.value).strip() for c in row if c.value is not None and str(c.value).strip()]
            if vals:
                rows.append(" | ".join(vals))
        if rows:
            chunks.append(f"=== SHEET: {name} ===\n" + "\n".join(rows))
    return "\n\n".join(chunks)
